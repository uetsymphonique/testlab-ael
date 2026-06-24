#!/usr/bin/env python3
"""
enrich_logsources.py — Enrich a plan CSV with log sources from mitre-knowledge-base/techniques.

Reads a CSV produced by export_ref_tables.py or assign-acw (same core columns, ACW column optional).
For each unique Technique ID in the CSV, looks up detection entries for the specified platform
in the MITRE knowledge-base technique files, then outputs one row per log source entry.

Output formats:
  - CSV (default, --out out.csv or stdout)
  - Excel (--excel or --out out.xlsx) with color-banded behavior groups for readability

Usage:
    python enrich_logsources.py plan.csv [--platform windows] [--techniques-dir PATH] [--out out.csv]
    python enrich_logsources.py plan.csv --excel --out out.xlsx
"""

import argparse
import csv
import re
import sys
from pathlib import Path

DEFAULT_TECHNIQUES_DIR = (
    Path(__file__).resolve().parent.parent.parent / "mitre-knowledge-base" / "techniques"
)
DEFAULT_PLATFORM = "windows"

ENRICH_HEADERS = [
    "Detection ID",
    "Detection Description",
    "Log Source",
    "Event Filter",
    "Data Component",
]
EMPTY_ENRICH = {h: "" for h in ENRICH_HEADERS}

_TECHNIQUE_HEADER_RE = re.compile(r"^### (T[\d.]+)\s*-\s*(.+)")
_DETECTION_ENTRY_RE = re.compile(r"^- \[(AN\d+)\] \*\*\[([^\]]+)\]\*\* (.+)")
_LOG_SOURCE_LINE_RE = re.compile(r"^\s+- \*\*Log sources:\*\* (.+)")
_LOG_SOURCE_ENTRY_RE = re.compile(r"`([^`]+)`\s*\(([^)]*)\)\s*\[([^\]]+)\]")


# ── Excel styling constants ──────────────────────────────────────────────────
# 4-color scheme: Calibrated × parity, Not Calibrated × parity
# Calibrated — green family
CAL_EVEN = "E2EFDA"   # light green
CAL_ODD  = "C5E0B4"   # darker green
# Not Calibrated — orange/peach family
NCAL_EVEN = "FCE4D6"  # light orange
NCAL_ODD  = "F8CBAD"  # darker orange
HEADER_FILL = "2F5496"  # dark blue
HEADER_FONT_COLOR = "FFFFFF"  # white
BORDER_COLOR = "B0B0B0"  # light gray gridlines


def read_csv_rows(csv_path: Path) -> tuple[list[str], list[dict]]:
    """Return (fieldnames, rows) from the plan CSV."""
    with csv_path.open(encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        rows = list(reader)
        fieldnames = list(reader.fieldnames or [])
    return fieldnames, rows


def parse_log_source_string(log_source_str: str) -> list[tuple[str, str, str]]:
    """Parse a log sources string into (source, event_filter, data_component) tuples."""
    return _LOG_SOURCE_ENTRY_RE.findall(log_source_str)


def extract_from_file(
    content: str,
    target_tids: set[str],
    platform_lower: str,
    seen_det_ids: set[tuple[str, str]],
) -> dict[str, list[dict]]:
    """
    Extract enrichment entries per technique ID from a tactic markdown file.
    Returns {tid: [enrich_dict, ...]} where each enrich_dict has ENRICH_HEADERS keys.
    `seen_det_ids` is mutated to deduplicate (tid, det_id) across tactic files.
    """
    lines = content.splitlines()
    results: dict[str, list[dict]] = {}

    current_tid: str | None = None
    in_detection = False
    pending_detection: tuple[str, str, str] | None = None  # (det_id, platform_label, desc)

    for line in lines:
        # ── Technique header ────────────────────────────────────────────────
        m = _TECHNIQUE_HEADER_RE.match(line)
        if m:
            current_tid = m.group(1).strip()
            in_detection = False
            pending_detection = None
            if current_tid not in target_tids:
                current_tid = None
            continue

        if current_tid is None:
            continue

        stripped = line.strip()

        # ── Section boundaries ───────────────────────────────────────────────
        if stripped == "**Detection**":
            in_detection = True
            continue

        if stripped == "**Procedure Examples**":
            in_detection = False
            pending_detection = None
            continue

        if not in_detection:
            continue

        # ── Detection entry bullet ───────────────────────────────────────────
        m = _DETECTION_ENTRY_RE.match(line)
        if m:
            pending_detection = (m.group(1), m.group(2), m.group(3).strip())
            continue

        # ── Log sources sub-bullet ────────────────────────────────────────────
        m = _LOG_SOURCE_LINE_RE.match(line)
        if m and pending_detection:
            det_id, det_platform_label, det_desc = pending_detection
            pending_detection = None

            if det_platform_label.lower() != platform_lower:
                continue

            dedup_key = (current_tid, det_id)
            if dedup_key in seen_det_ids:
                continue
            seen_det_ids.add(dedup_key)

            log_sources = parse_log_source_string(m.group(1))
            entries = results.setdefault(current_tid, [])

            if log_sources:
                for source, event_filter, data_component in log_sources:
                    entries.append(
                        {
                            "Detection ID": det_id,
                            "Detection Description": det_desc,
                            "Log Source": source,
                            "Event Filter": event_filter,
                            "Data Component": data_component,
                        }
                    )
            else:
                entries.append(
                    {
                        "Detection ID": det_id,
                        "Detection Description": det_desc,
                        "Log Source": "",
                        "Event Filter": "",
                        "Data Component": "",
                    }
                )

    return results


def lookup_logsources(
    techniques_dir: Path,
    target_tids: set[str],
    platform: str,
) -> tuple[dict[str, list[dict]], set[str], set[str]]:
    """
    Returns (enrich_map, tids_not_in_kb, tids_no_platform_coverage).

    - enrich_map: {tid: [enrich_dict, ...]} — one entry per log source
    - tids_not_in_kb: technique IDs not found in any techniques file
    - tids_no_platform_coverage: found in KB but no detection entry for the requested platform
    """
    platform_lower = platform.lower()

    tids_in_kb: set[str] = set()
    enrich_map: dict[str, list[dict]] = {}
    seen_det_ids: set[tuple[str, str]] = set()

    for tactic_file in sorted(techniques_dir.glob("*.md")):
        content = tactic_file.read_text(encoding="utf-8")

        for m in re.finditer(r"^### (T[\d.]+)\s*-", content, re.MULTILINE):
            if m.group(1) in target_tids:
                tids_in_kb.add(m.group(1))

        file_map = extract_from_file(content, target_tids, platform_lower, seen_det_ids)
        for tid, entries in file_map.items():
            enrich_map.setdefault(tid, []).extend(entries)

    tids_not_in_kb = target_tids - tids_in_kb
    tids_no_platform_coverage = tids_in_kb - enrich_map.keys()

    return enrich_map, tids_not_in_kb, tids_no_platform_coverage


def build_output_rows(
    input_rows: list[dict],
    enrich_map: dict[str, list[dict]],
) -> tuple[list[dict], list[int]]:
    """
    Build expanded output rows with enrichment, tracking group boundaries.

    Returns (out_rows, group_ids) where:
      - out_rows: list of merged dicts (input columns + enrichment columns)
      - group_ids: parallel list — output rows with the same group_id belong to the
        same original input row (same behavior). Used for Excel color banding.
    """
    out_rows: list[dict] = []
    group_ids: list[int] = []

    for idx, row in enumerate(input_rows):
        tid = row.get("Technique ID", "").strip()
        enrichments = enrich_map.get(tid, [])
        if enrichments:
            for enrich in enrichments:
                out_rows.append({**row, **enrich})
                group_ids.append(idx)
        else:
            out_rows.append({**row, **EMPTY_ENRICH})
            group_ids.append(idx)

    return out_rows, group_ids


def write_csv(
    out_path: Path | None,
    fieldnames: list[str],
    out_rows: list[dict],
) -> None:
    """Write output rows as CSV (to file or stdout)."""
    if out_path:
        out_path.parent.mkdir(parents=True, exist_ok=True)
        with out_path.open("w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(out_rows)
        print(f"[+] Wrote {len(out_rows)} row(s) to {out_path}", file=sys.stderr)
    else:
        writer = csv.DictWriter(sys.stdout, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(out_rows)


def write_excel(
    out_path: Path,
    fieldnames: list[str],
    out_rows: list[dict],
    group_ids: list[int],
) -> None:
    """Write output rows as a color-banded Excel workbook."""
    from openpyxl import Workbook
    from openpyxl.styles import (
        Alignment,
        Border,
        Font,
        NamedStyle,
        PatternFill,
        Side,
    )
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Enriched Plan"

    # ── Styles ────────────────────────────────────────────────────────────────
    header_fill = PatternFill(start_color=HEADER_FILL, end_color=HEADER_FILL, fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color=HEADER_FONT_COLOR)
    thin_border = Border(
        left=Side(style="thin", color=BORDER_COLOR),
        right=Side(style="thin", color=BORDER_COLOR),
        top=Side(style="thin", color=BORDER_COLOR),
        bottom=Side(style="thin", color=BORDER_COLOR),
    )
    wrap_alignment = Alignment(wrap_text=True, vertical="top")

    # Pre-build PatternFills for each of the 4 color bands
    fill_cal_even = PatternFill(start_color=CAL_EVEN, end_color=CAL_EVEN, fill_type="solid")
    fill_cal_odd  = PatternFill(start_color=CAL_ODD,  end_color=CAL_ODD,  fill_type="solid")
    fill_ncal_even = PatternFill(start_color=NCAL_EVEN, end_color=NCAL_EVEN, fill_type="solid")
    fill_ncal_odd  = PatternFill(start_color=NCAL_ODD,  end_color=NCAL_ODD,  fill_type="solid")

    # Find where original columns end and enrichment columns begin
    enrich_start_col = len(fieldnames) - len(ENRICH_HEADERS) + 1  # 1-based

    # ── Header row ────────────────────────────────────────────────────────────
    for col_idx, header in enumerate(fieldnames, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = thin_border

    ws.row_dimensions[1].height = 28

    # ── Data rows: 4-color banding (Calibrated × parity, Not Calibrated × parity)
    # Each unique group_id inherits the Category of its first output row.
    # Within each Category class, groups alternate colors independently.
    unique_gids = sorted(set(group_ids))
    gid_category: dict[int, str] = {}
    for gid in unique_gids:
        # First output row for this group_id
        first = next(r for r, g in zip(out_rows, group_ids) if g == gid)
        gid_category[gid] = first.get("Category", "").strip()

    # Counters per category to alternate within each class
    cat_counter: dict[str, int] = {}
    group_band_map: dict[int, PatternFill] = {}
    for gid in unique_gids:
        cat = gid_category[gid]
        idx = cat_counter.get(cat, 0)
        cat_counter[cat] = idx + 1
        is_cal = cat.startswith("Calibrated")
        if is_cal:
            group_band_map[gid] = fill_cal_even if idx % 2 == 0 else fill_cal_odd
        else:
            group_band_map[gid] = fill_ncal_even if idx % 2 == 0 else fill_ncal_odd

    for row_idx, (row, gid) in enumerate(zip(out_rows, group_ids)):
        excel_row = row_idx + 2  # 1-based, row 1 is header
        fill = group_band_map[gid]

        for col_idx, field in enumerate(fieldnames, start=1):
            value = row.get(field, "")
            cell = ws.cell(row=excel_row, column=col_idx, value=value)
            cell.fill = fill
            cell.border = thin_border
            cell.alignment = wrap_alignment

            # Bold the enrichment columns to visually separate from original data
            if col_idx >= enrich_start_col:
                cell.font = Font(name="Calibri", size=10, bold=True)

    # ── Column widths: auto-fit ───────────────────────────────────────────────
    for col_idx, field in enumerate(fieldnames, start=1):
        # Determine max content width in this column
        max_len = len(str(field))
        for row in out_rows:
            val = str(row.get(field, ""))
            # For wrapped cells, use the longest single line
            lines = val.split("\n")
            line_max = max((len(line) for line in lines), default=0)
            max_len = max(max_len, line_max)

        # Clamp to reasonable bounds
        col_width = min(max_len + 3, 60)
        col_width = max(col_width, 10)
        ws.column_dimensions[get_column_letter(col_idx)].width = col_width

    # ── Freeze header row ────────────────────────────────────────────────────
    ws.freeze_panes = "A2"

    # ── Auto-filter on header ────────────────────────────────────────────────
    ws.auto_filter.ref = f"A1:{get_column_letter(len(fieldnames))}{len(out_rows) + 1}"

    # ── Add a legend sheet ───────────────────────────────────────────────────
    ws2 = wb.create_sheet("Legend")
    ws2.column_dimensions["A"].width = 26
    ws2.column_dimensions["B"].width = 55

    ws2.cell(row=1, column=1, value="Color").font = Font(bold=True)
    ws2.cell(row=1, column=2, value="Meaning").font = Font(bold=True)

    ws2.cell(row=2, column=1, value="Calibrated — Even group").fill = fill_cal_even
    ws2.cell(row=2, column=2, value="Calibrated behavior, even-numbered group within Calibrated class → counts toward detection-rate denominator")

    ws2.cell(row=3, column=1, value="Calibrated — Odd group").fill = fill_cal_odd
    ws2.cell(row=3, column=2, value="Calibrated behavior, odd-numbered group within Calibrated class — same semantic, alternating shade to separate adjacent groups")

    ws2.cell(row=4, column=1, value="Not Calibrated — Even group").fill = fill_ncal_even
    ws2.cell(row=4, column=2, value="Not Calibrated behavior, even-numbered group within Not Calibrated class → excluded from detection-rate denominator")

    ws2.cell(row=5, column=1, value="Not Calibrated — Odd group").fill = fill_ncal_odd
    ws2.cell(row=5, column=2, value="Not Calibrated behavior, odd-numbered group within Not Calibrated class — same semantic, alternating shade")

    ws2.cell(row=7, column=1, value="Bold columns").font = Font(bold=True)
    ws2.cell(
        row=7, column=2,
        value="Enrichment columns (Detection ID through Data Component) — added by enrich_logsources.py",
    )

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_path))
    print(f"[+] Wrote {len(out_rows)} row(s) to {out_path} (Excel with color bands)", file=sys.stderr)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Enrich a plan CSV (from export_ref_tables.py or assign-acw) with log sources "
            "from mitre-knowledge-base/techniques for a given platform."
        )
    )
    parser.add_argument("csv_file", type=Path, help="Input plan CSV file.")
    parser.add_argument(
        "--platform",
        default=DEFAULT_PLATFORM,
        help=f"Detection platform label to filter on (case-insensitive). Default: {DEFAULT_PLATFORM}",
    )
    parser.add_argument(
        "--techniques-dir",
        type=Path,
        default=DEFAULT_TECHNIQUES_DIR,
        help=f"Path to mitre-knowledge-base/techniques directory. Default: {DEFAULT_TECHNIQUES_DIR}",
    )
    parser.add_argument(
        "--out",
        type=Path,
        default=None,
        help="Output path. Defaults to stdout (CSV). Use .xlsx extension or --excel for Excel output.",
    )
    parser.add_argument(
        "--excel",
        action="store_true",
        default=False,
        help="Force Excel (.xlsx) output. Also auto-detected from --out extension.",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()

    if not args.csv_file.exists():
        print(f"[!] CSV file not found: {args.csv_file}", file=sys.stderr)
        return 1

    if not args.techniques_dir.is_dir():
        print(f"[!] Techniques directory not found: {args.techniques_dir}", file=sys.stderr)
        return 1

    input_fieldnames, input_rows = read_csv_rows(args.csv_file)
    if not input_rows:
        print("[!] No rows found in input CSV.", file=sys.stderr)
        return 1

    target_tids = {r.get("Technique ID", "").strip() for r in input_rows} - {""}
    print(
        f"[*] {len(target_tids)} unique technique(s) in plan CSV, platform filter: [{args.platform}]",
        file=sys.stderr,
    )

    enrich_map, not_in_kb, no_coverage = lookup_logsources(
        args.techniques_dir, target_tids, args.platform
    )

    if not_in_kb:
        for tid in sorted(not_in_kb):
            print(f"[!] Not found in knowledge base: {tid}", file=sys.stderr)

    if no_coverage:
        for tid in sorted(no_coverage):
            print(f"[-] No [{args.platform}] detection entry: {tid}", file=sys.stderr)

    output_fieldnames = input_fieldnames + ENRICH_HEADERS
    out_rows, group_ids = build_output_rows(input_rows, enrich_map)

    total_enrich = sum(len(v) for v in enrich_map.values())
    print(
        f"[+] {len(out_rows)} output row(s) ({total_enrich} log source entries across {len(enrich_map)} technique(s))",
        file=sys.stderr,
    )

    # ── Decide output format ──────────────────────────────────────────────────
    use_excel = args.excel
    out_path = args.out
    if out_path and out_path.suffix.lower() == ".xlsx":
        use_excel = True

    if use_excel:
        if not out_path:
            out_path = args.csv_file.with_suffix(".xlsx")
        write_excel(out_path, output_fieldnames, out_rows, group_ids)
    else:
        write_csv(out_path, output_fieldnames, out_rows)

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
