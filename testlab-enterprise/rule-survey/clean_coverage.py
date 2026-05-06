#!/usr/bin/env python3
"""
clean_coverage.py - Deduplicate technique IDs and append technique names to coverage.tsv.
Outputs: coverage_clean.tsv (UTF-8 BOM, tab-separated)
         coverage_clean.xlsx (Excel, with header formatting)

Requires: pip install openpyxl
"""

import csv
import os
import re

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    HAVE_OPENPYXL = True
except ImportError:
    HAVE_OPENPYXL = False

COVERAGE_TSV = os.path.join(os.path.dirname(os.path.abspath(__file__)), "coverage.tsv")
TECHNIQUE_NAMES_TXT = r"d:\vcs\ael\mitre-knowledge-base\technique_names.txt"
OUTPUT_TSV = os.path.join(os.path.dirname(os.path.abspath(__file__)), "coverage_clean.tsv")
OUTPUT_XLSX = os.path.join(os.path.dirname(os.path.abspath(__file__)), "coverage_clean.xlsx")


def load_technique_names(path):
    """Parse technique_names.txt (array format) into {TID: full_name}."""
    tid_to_name = {}
    pattern = re.compile(r'^"?(T\d+(?:\.\d+)?)\s+-\s+(.+?)"?,?\s*$')
    with open(path, encoding="utf-8-sig") as f:
        for line in f:
            m = pattern.match(line.strip())
            if m:
                tid, name = m.group(1), m.group(2)
                tid_to_name[tid] = name
    return tid_to_name


def dedup_ordered(tid_str):
    """Deduplicate comma-separated TIDs, preserving first-occurrence order."""
    if not tid_str or not tid_str.strip():
        return []
    seen = set()
    result = []
    for tid in tid_str.split(","):
        tid = tid.strip()
        if tid and tid not in seen:
            seen.add(tid)
            result.append(tid)
    return result


def main():
    tid_to_name = load_technique_names(TECHNIQUE_NAMES_TXT)
    print(f"Technique names loaded: {len(tid_to_name)}")

    rows = []
    with open(COVERAGE_TSV, encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f, delimiter="\t")
        header = next(reader)
        for row in reader:
            # Skip trailing marker lines (e.g. a lone ".")
            if row and row[0].strip() in (".", ""):
                continue
            rows.append(row)

    try:
        tech_col = header.index("Technique")
    except ValueError:
        print("ERROR: 'Technique' column not found in header.")
        return

    # Insert "Technique Names" column right after "Technique"
    new_header = header[: tech_col + 1] + ["Technique Names"] + header[tech_col + 1 :]

    changed = 0
    new_rows = []
    for row in rows:
        # Pad short rows
        while len(row) <= tech_col:
            row.append("")

        raw = row[tech_col]
        tids = dedup_ordered(raw)
        deduped = ",".join(tids)
        names = "; ".join(tid_to_name.get(t, f"[UNKNOWN: {t}]") for t in tids)

        if deduped != raw:
            changed += 1

        new_row = row[:tech_col] + [deduped] + [names] + row[tech_col + 1 :]
        new_rows.append(new_row)

    with open(OUTPUT_TSV, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f, delimiter="\t")
        writer.writerow(new_header)
        writer.writerows(new_rows)

    print(f"Rows processed   : {len(new_rows)}")
    print(f"Rows deduped     : {changed}")
    print(f"Output written to: {OUTPUT_TSV}")

    if HAVE_OPENPYXL:
        _write_xlsx(new_header, new_rows)
    else:
        print("WARNING: openpyxl not installed — skipping XLSX output. Run: pip install openpyxl")


def _write_xlsx(header, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Coverage"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(fill_type="solid", fgColor="2F5597")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.append(header)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    for row in rows:
        ws.append(row)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # Column widths (approximate)
    col_widths = {
        "A": 6,   # STT
        "B": 45,  # Rule name
        "C": 30,  # Tactic
        "D": 35,  # Technique
        "E": 60,  # Technique Names
        "F": 10,  # EventID
        "G": 25,  # Data Component
        "H": 12,  # Release Level
        "I": 10,  # Rule Status
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    wb.save(OUTPUT_XLSX)
    print(f"Excel written to : {OUTPUT_XLSX}")


if __name__ == "__main__":
    main()
